import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET

SUPPORTED_EXTENSIONS = [
    ".pdf", ".doc", ".docx", ".ppt", ".pptx",
    ".xls", ".xlsx", ".csv", ".txt", ".md", ".html", ".htm",
]

MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15MB, matches frontend spec

# Legacy OLE Compound File signature (old binary .doc/.xls/.ppt formats,
# pre-2007). These aren't parseable with python-docx/zipfile, so they're
# detected up front and rejected with a clear message instead of silently
# falling back to garbage-decoded bytes.
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _is_legacy_ole_binary(content_bytes: bytes) -> bool:
    return content_bytes[:8] == _OLE_SIGNATURE


def extract_text_from_file(filename: str, content_bytes: bytes) -> str:
    """
    Extract text content from uploaded document bytes based on file extension.
    """
    ext = os.path.splitext(filename)[1].lower() if "." in filename else ""

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext or 'unknown'}'. Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
        )

    if len(content_bytes) > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File size ({len(content_bytes)} bytes) exceeds maximum limit of 15MB."
        )

    if ext == ".pdf":
        text = _extract_pdf(content_bytes)
    elif ext == ".docx":
        text = _extract_docx(content_bytes)
    elif ext == ".doc":
        text = _extract_doc(content_bytes)
    elif ext in (".pptx", ".ppt"):
        text = _extract_pptx(content_bytes)
    elif ext in (".xlsx", ".xls"):
        text = _extract_xlsx(content_bytes)
    elif ext in (".html", ".htm"):
        text = _extract_html(content_bytes)
    else:
        text = _decode_text(content_bytes)

    cleaned = text.strip()
    if not cleaned:
        raise ValueError("The uploaded document contains no readable text.")

    return cleaned


def _decode_text(content_bytes: bytes) -> str:
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            text = content_bytes.decode(encoding)
            return text.strip()
        except UnicodeDecodeError:
            continue
    return content_bytes.decode("utf-8", errors="ignore").strip()


def _extract_pdf(content_bytes: bytes) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=content_bytes, filetype="pdf")
        pages_text = []
        for page in doc:
            text = page.get_text()
            if text.strip():
                pages_text.append(text.strip())
        doc.close()
        return "\n\n".join(pages_text)
    except Exception as e:
        raise ValueError(f"Could not parse PDF document: {e}")


def _extract_docx(content_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(content_bytes))
        lines = []
        for para in doc.paragraphs:
            if para.text.strip():
                lines.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    lines.append(" | ".join(row_text))
        return "\n".join(lines)
    except Exception as e:
        raise ValueError(f"Could not parse DOCX document: {e}")


def _extract_doc(content_bytes: bytes) -> str:
    if _is_legacy_ole_binary(content_bytes):
        raise ValueError(
            "This looks like a legacy binary .doc file (Word 97-2003), "
            "which isn't supported. Please re-save it as .docx and try again."
        )

    try:
        import docx
        doc = docx.Document(io.BytesIO(content_bytes))
        lines = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        if lines:
            return "\n".join(lines)
    except Exception:
        pass

    raw_text = _decode_text(content_bytes)
    printable = re.findall(r"[\x20-\x7E\t\r\n]{4,}", raw_text)
    return "\n".join(printable)


def _extract_pptx(content_bytes: bytes) -> str:
    if _is_legacy_ole_binary(content_bytes):
        raise ValueError(
            "This looks like a legacy binary .ppt file (PowerPoint 97-2003), "
            "which isn't supported. Please re-save it as .pptx and try again."
        )

    try:
        import pptx
        prs = pptx.Presentation(io.BytesIO(content_bytes))
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_runs.append(shape.text.strip())
        if text_runs:
            return "\n".join(text_runs)
    except Exception:
        pass

    try:
        zf = zipfile.ZipFile(io.BytesIO(content_bytes))
        texts = []
        for name in sorted(zf.namelist()):
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                xml_content = zf.read(name)
                tree = ET.fromstring(xml_content)
                slide_text = []
                for elem in tree.iter():
                    if elem.tag.endswith("}t") and elem.text:
                        slide_text.append(elem.text.strip())
                if slide_text:
                    texts.append(" ".join(slide_text))
        if texts:
            return "\n\n".join(texts)
    except Exception:
        pass

    return _decode_text(content_bytes)


def _extract_xlsx(content_bytes: bytes) -> str:
    if _is_legacy_ole_binary(content_bytes):
        raise ValueError(
            "This looks like a legacy binary .xls file (Excel 97-2003), "
            "which isn't supported. Please re-save it as .xlsx and try again."
        )

    try:
        zf = zipfile.ZipFile(io.BytesIO(content_bytes))
        shared_strings = []
        if "xl/sharedStrings.xml" in zf.namelist():
            tree = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for elem in tree.iter():
                if elem.tag.endswith("}t") and elem.text:
                    shared_strings.append(elem.text.strip())
        if shared_strings:
            return "\n".join(shared_strings)
    except Exception:
        pass

    return _decode_text(content_bytes)


def _extract_html(content_bytes: bytes) -> str:
    html_str = _decode_text(content_bytes)
    html_str = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html_str, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", html_str)
    text = re.sub(r"\s+", " ", text)
    return text.strip()
