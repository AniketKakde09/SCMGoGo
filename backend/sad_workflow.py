"""SAD proposal generation. Never writes to the source workbook or Jira."""
from __future__ import annotations

import io
import json
import logging
import os
import time
import re
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
from fastapi import HTTPException

from llm import LLMClient
from security import sanitize_text, sanitize_dataframe


MAX_BYTES = 8 * 1024 * 1024
MAX_CHARS = 70000
TYPES = ('Epic', 'Feature', 'Story', 'Task')


def extract_document(filename: str, content: bytes) -> str:
    ext = Path(filename or '').suffix.lower()

    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'SAD must be 8 MB or smaller')

    if ext in ('.txt', '.md'):
        try:
            return content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise HTTPException(422, 'Upload UTF-8 text')

    if ext == '.docx':
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))

            return (
                '\n'.join(
                    p.text
                    for p in doc.paragraphs
                    if p.text.strip()
                )
                + '\n'
                + '\n'.join(
                    ' | '.join(c.text for c in row.cells)
                    for table in doc.tables
                    for row in table.rows
                )
            )

        except Exception as exc:
            raise HTTPException(
                422,
                f'Unable to read DOCX: {exc}'
            )

    if ext == '.pdf':
        try:
            from pypdf import PdfReader

            return '\n'.join(
                page.extract_text() or ''
                for page in PdfReader(io.BytesIO(content)).pages
            )

        except Exception as exc:
            raise HTTPException(
                422,
                f'Unable to read PDF: {exc}'
            )

    raise HTTPException(
        415,
        'Supported SAD formats: PDF, DOCX, TXT, MD'
    )


def sections_from_text(text: str):
    sections = []
    current = None

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        heading = re.match(
            r'^(?:#{1,4}\s+|(?:section\s+)?\d+(?:\.\d+)*[.)]?\s+)(.{3,100})$',
            line,
            re.I
        )

        if heading:
            current = {
                'id': f'section-{len(sections) + 1}',
                'title': heading.group(1).strip(),
                'text': ''
            }
            sections.append(current)

        elif current:
            current['text'] += line + '\n'

        else:
            current = {
                'id': 'section-1',
                'title': 'Overview',
                'text': line + '\n'
            }
            sections.append(current)

    if len(sections) > 80:
        raise HTTPException(
            413,
            'SAD has more than 80 sections; split it into separate documents'
        )

    return sections


def _tokens(text):
    return set(
        re.findall(
            r'[a-z0-9]{3,}',
            str(text).lower()
        )
    ) - {
        'the',
        'and',
        'for',
        'with',
        'from',
        'this',
        'that',
        'into',
        'should',
        'shall',
        'must',
        'will',
        'have',
        'using'
    }


def _candidate(proposal, backlog):
    """Conservative lexical screening, NOT a semantic duplicate verdict.

    Only compare actionable work of the same type. Shared domain vocabulary,
    matching architecture sections, and parent/child relationships are not duplicates.
    """
    kind = str(
        proposal.get('type', '')
    ).strip().lower()

    if kind not in ('story', 'task'):
        return []

    title = str(
        proposal.get('title', '')
    ).strip()

    description = str(
        proposal.get('description', '')
    ).strip()

    if not title or not description:
        return []

    title_tokens = _tokens(title)

    normalized_title = re.sub(
        r'[^a-z0-9]+',
        ' ',
        title.casefold()
    ).strip()

    description_tokens = _tokens(description)

    if len(title_tokens) < 2 or len(description_tokens) < 3:
        return []

    found = []

    for _, row in backlog.iterrows():
        if str(
            row.get('Type', '')
        ).strip().lower() != kind:
            continue

        existing_title = str(
            row.get('Title', '')
        ).strip()

        existing_description = str(
            row.get('Description', '')
        ).strip()

        if not existing_title or not existing_description:
            continue

        other_title_tokens = _tokens(existing_title)
        other_description_tokens = _tokens(existing_description)

        if not other_title_tokens or not other_description_tokens:
            continue

        title_overlap = (
            len(title_tokens & other_title_tokens)
            / max(
                1,
                len(title_tokens | other_title_tokens)
            )
        )

        body_overlap = (
            len(description_tokens & other_description_tokens)
            / max(
                1,
                len(description_tokens | other_description_tokens)
            )
        )

        title_similarity = SequenceMatcher(
            None,
            title.casefold(),
            existing_title.casefold()
        ).ratio()

        normalized_existing = re.sub(
            r'[^a-z0-9]+',
            ' ',
            existing_title.casefold()
        ).strip()

        exact_title = (
            normalized_title == normalized_existing
        )

        # Exact titles are strong evidence; near-identical titles also need
        # substantial scope overlap. Related architecture vocabulary alone is not enough.
        if (
            (exact_title and body_overlap >= .35)
            or (
                title_similarity >= .91
                and body_overlap >= .60
            )
            or (
                title_overlap >= .86
                and body_overlap >= .78
            )
        ):
            score = round(
                max(
                    title_similarity,
                    title_overlap
                ) * .55
                + body_overlap * .45,
                3
            )

            found.append({
                'ticket_id': str(
                    row.get('TicketID', '')
                ).strip(),
                'title': existing_title,
                'type': str(
                    row.get('Type', '')
                ),
                'status': str(
                    row.get('Status', '')
                ),
                'score': score,
                'confidence': 'high',
                'reason': (
                    'Highly similar title and overlapping description; '
                    'verify acceptance criteria and scope.'
                )
            })

    return sorted(
        (
            item
            for item in found
            if item['ticket_id']
        ),
        key=lambda item: item['score'],
        reverse=True
    )[:3]


logger = logging.getLogger(__name__)

# Character budgets are conservative approximations, not model token guarantees.
CHUNK_CHARS = 2600
MAX_CHUNKS = 60
MAX_TICKETS = 60
BATCH_OUTPUT_TOKENS = 2400
MAX_SPLIT_DEPTH = 4


def _chunks(sections):
    """Never silently truncate SAD content; retain each source section ID."""
    chunks = []

    for section in sections:
        content = section['text'].strip()

        if not content:
            continue

        # Split on word boundaries, including when a single paragraph is huge.
        words = content.split()
        part = []
        size = 0

        for word in words:
            if len(word) > CHUNK_CHARS:
                raise HTTPException(
                    422,
                    f"Section {section['id']} contains an oversized unbroken token"
                )

            if part and size + len(word) + 1 > CHUNK_CHARS:
                chunks.append({
                    'id': section['id'],
                    'title': section['title'],
                    'text': ' '.join(part)
                })

                part, size = [], 0

            part.append(word)
            size += len(word) + 1

        if part:
            chunks.append({
                'id': section['id'],
                'title': section['title'],
                'text': ' '.join(part)
            })

    if len(chunks) > MAX_CHUNKS:
        raise HTTPException(
            413,
            f'SAD requires {len(chunks)} batches '
            f'(maximum {MAX_CHUNKS}); process as separate documents'
        )

    return chunks


def _source_excerpt(chunk, suggested):
    """Use text owned by the document, never an unverified model quotation."""
    source = chunk['text']

    suggested = str(
        suggested or ''
    ).strip()

    if suggested and suggested in source:
        return suggested[:500]

    # Whitespace differences are common in LLM quotations. Find a source-owned
    # excerpt if possible; otherwise use a short literal excerpt from this chunk.
    if suggested:
        normalized = re.sub(
            r'\s+',
            ' ',
            suggested
        ).strip().lower()

        words = source.split()

        for start in range(len(words)):
            candidate = ' '.join(
                words[
                    start:start + min(
                        24,
                        len(words) - start
                    )
                ]
            )

            if normalized in candidate.lower():
                return candidate[:500]

    return source[:min(240, len(source))]


def _validate_batch(tickets, chunk):
    if not isinstance(tickets, list) or len(tickets) > 3:
        raise ValueError(
            'Expected tickets array containing at most 3 items'
        )

    by_id = {}

    for item in tickets:
        if (
            not isinstance(item, dict)
            or item.get('type') not in TYPES
            or not str(item.get('title', '')).strip()
        ):
            raise ValueError(
                'Ticket has invalid type or title'
            )

        dod = item.get(
            'definition_of_done',
            []
        )

        if dod is None:
            dod = []

        if not isinstance(dod, list):
            raise ValueError(
                'definition_of_done must be a list'
            )

        if len(dod) > 2:
            raise ValueError(
                'definition_of_done must contain at most 2 items'
            )

        if item.get('type') != 'Story':
            dod = []

        for entry in dod:
            if not str(entry).strip():
                raise ValueError(
                    'definition_of_done cannot contain blank items'
                )

        item['definition_of_done'] = [
            str(x).strip()[:300]
            for x in dod
            if str(x).strip()
        ][:2]

        key = str(
            item.get('temp_id', '')
        ).strip()

        if not key or key in by_id:
            raise ValueError(
                'Missing or duplicate ticket temp_id'
            )

        by_id[key] = item

        # The chunk ID is authoritative; don't allow invented source references.
        if item.get('source_section_id') != chunk['id']:
            raise ValueError(
                'Ticket source section does not match input'
            )

        item['source_excerpt'] = _source_excerpt(
            chunk,
            item.get('source_excerpt')
        )

    expected = {
        'Feature': 'Epic',
        'Story': 'Feature',
        'Task': 'Feature'
    }

    for item in tickets:
        parent_id = str(
            item.get('parent_temp_id') or ''
        ).strip()

        parent = by_id.get(parent_id)

        if item['type'] == 'Epic':
            if parent_id:
                raise ValueError(
                    'Epic must not have a parent'
                )

        elif (
            not parent
            or parent['type'] != expected[item['type']]
        ):
            raise ValueError(
                'Ticket has invalid parent hierarchy'
            )

    return tickets


def _split_chunk(chunk):
    """Split a failed chunk without losing or changing source section identity."""
    words = chunk['text'].split()

    if len(words) < 24:
        return None

    midpoint = len(words) // 2

    left = ' '.join(words[:midpoint])
    right = ' '.join(words[midpoint:])

    if not left or not right:
        return None

    return [
        {
            **chunk,
            'text': left
        },
        {
            **chunk,
            'text': right
        }
    ]


def _generate_batch(llm, chunk, index, total, depth=0):
    """
    Generate a very small, self-contained hierarchy for one SAD chunk.

    The model may return at most 3 tickets TOTAL.

    Valid structures are only:

        Epic
        └── Feature
            └── Story

    or:

        Epic
        └── Feature
            └── Task

    Every parent referenced by a ticket MUST exist in the same response.
    """

    base_prompt = (
        'Convert ONLY the supplied SAD excerpt into a SMALL draft delivery hierarchy. '

        'IMPORTANT HARD LIMIT: return AT MOST 3 tickets TOTAL. '
        'Never return 4 or more tickets. '
        'Do not create one ticket for every requirement. '
        'Select the smallest useful hierarchy that represents the supplied excerpt. '

        'The hierarchy MUST be self-contained within this response. '

        'Allowed hierarchy structures are ONLY: '
        '(1) Epic -> Feature -> Story, or '
        '(2) Epic -> Feature -> Task. '

        'An Epic has no parent. '
        'A Feature MUST have an Epic as its parent. '
        'A Story MUST have a Feature as its parent. '
        'A Task MUST have a Feature as its parent. '

        'Every parent_temp_id MUST exactly match the temp_id of another ticket '
        'in the same returned tickets array. '
        'Never reference a parent that is not returned. '

        'If you cannot create a valid hierarchy from the excerpt, return '
        '{"tickets":[]} instead of inventing relationships. '

        'Return ONLY a valid JSON object with root key "tickets". '
        'Do not wrap JSON in markdown. '

        'Each ticket must contain: '
        'temp_id, parent_temp_id, type, title, description, '
        'acceptance_criteria, definition_of_done, source_section_id, '
        'source_excerpt, assumptions. '

        'Use exactly these ID patterns when applicable: '
        'Epic=e1, Feature=f1, Story=s1, Task=t1. '

        'Valid examples: '
        '{"temp_id":"e1","parent_temp_id":"","type":"Epic"}, '
        '{"temp_id":"f1","parent_temp_id":"e1","type":"Feature"}, '
        '{"temp_id":"s1","parent_temp_id":"f1","type":"Story"}. '

        'OR: '
        '{"temp_id":"e1","parent_temp_id":"","type":"Epic"}, '
        '{"temp_id":"f1","parent_temp_id":"e1","type":"Feature"}, '
        '{"temp_id":"t1","parent_temp_id":"f1","type":"Task"}. '

        'Do NOT return Story without Feature. '
        'Do NOT return Task without Feature. '
        'Do NOT return Feature without Epic. '

        'Capture necessary work such as infrastructure, functional work, '
        'QA, or security only when it is directly supported by this excerpt. '

        'description must be at most 120 characters. '
        'acceptance_criteria must contain at most 2 short strings. '
        'definition_of_done must contain at most 2 short strings and only '
        'when supported by the excerpt. '
        'DoD applies only to Stories; for Epic, Feature, and Task use []. '
        'assumptions must contain at most 1 short string. '

        'Use ONLY the provided source_section_id. '
        'Every ticket must trace directly to the supplied SAD excerpt. '

        'If there is no actionable requirement, return {"tickets":[]}. '

        'Keep the entire response short and complete.'
    )

    payload = json.dumps(
        {
            'source_section_id': chunk['id'],
            'section_title': chunk['title'],
            'excerpt': chunk['text']
        },
        ensure_ascii=False
    )

    last_error = None

    for attempt in range(2):
        try:
            prompt = base_prompt

            # On retry, explicitly tell the model what went wrong.
            if last_error:
                prompt += (
                    '\n\nPREVIOUS ATTEMPT WAS INVALID.\n'
                    f'Validation error: {str(last_error)[:500]}\n'
                    'Correct the hierarchy and return a completely valid '
                    'JSON response. Do not repeat the invalid structure.'
                )

            result = llm.generate_json(
                prompt,
                payload,
                max_tokens=BATCH_OUTPUT_TOKENS
            )

            if not isinstance(result, dict):
                raise ValueError(
                    'Model response must be a JSON object'
                )

            tickets = _validate_batch(
                result.get('tickets'),
                chunk
            )

            logger.info(
                'SAD batch %s/%s depth=%s generated %s tickets',
                index,
                total,
                depth,
                len(tickets)
            )

            return tickets

        except (
            ValueError,
            TypeError,
            KeyError,
            RuntimeError
        ) as exc:

            last_error = exc

            logger.warning(
                'SAD batch %s/%s depth=%s attempt=%s failed: %s: %.250s',
                index,
                total,
                depth,
                attempt + 1,
                type(exc).__name__,
                str(exc)
            )

            if attempt == 0:
                time.sleep(0.2)

    # A smaller input can recover from completion limits or
    # structurally difficult excerpts.
    halves = (
        _split_chunk(chunk)
        if depth < MAX_SPLIT_DEPTH
        else None
    )

    if halves:
        logger.info(
            'Splitting failed SAD batch %s/%s at depth %s',
            index,
            total,
            depth
        )

        combined = []

        for half in halves:
            half_tickets = _generate_batch(
                llm,
                half,
                index,
                total,
                depth + 1
            )

            # Each split is an independent local hierarchy.
            # Namespace IDs so two halves cannot collide.
            prefix = uuid.uuid4().hex[:8]

            for item in half_tickets:
                old_id = item['temp_id']

                item['temp_id'] = (
                    f'{prefix}-{old_id}'
                )

                if item.get('parent_temp_id'):
                    item['parent_temp_id'] = (
                        f'{prefix}-{item["parent_temp_id"]}'
                    )

            combined.extend(half_tickets)

        return combined

    raise HTTPException(
        502,
        f'SAD batch {index}/{total} failed after retry '
        f'and smaller-chunk recovery '
        f'({type(last_error).__name__}: {str(last_error)[:180]}). '
        'No partial proposal was returned.'
    ) from last_error


# The chunk model emits local scaffolding (Epic -> Feature -> leaf).
# Those scaffolds are NOT the document-wide hierarchy.
# Reconcile them before returning.


def _clean_title(value):
    return re.sub(
        r"\s+",
        " ",
        str(value or "")
    ).strip()[:255]


def _similarity(a, b):
    a = _clean_title(a).lower()
    b = _clean_title(b).lower()

    if not a or not b:
        return 0.0

    ta, tb = _tokens(a), _tokens(b)

    return max(
        SequenceMatcher(
            None,
            a,
            b
        ).ratio()
        if min(len(a), len(b)) >= 12
        else 0.0,
        len(ta & tb) / max(
            1,
            len(ta | tb)
        )
    )


def _reconcile_hierarchy(llm, staged):
    """Group local scaffolds across ALL chunks; preserve every generated leaf.

    Small classification calls use a compact global catalog. A failed classifier
    never discards work: conservative deterministic grouping is the fallback.
    """
    by_id = {
        item['id']: item
        for item in staged
    }

    leaves = [
        item
        for item in staged
        if item['type'] in ('Story', 'Task')
    ]

    # A batch can legitimately emit only an Epic or Feature.
    # Retain that work.
    for item in staged:

        if (
            item['type'] == 'Feature'
            and not any(
                x['parent_id'] == item['id']
                for x in leaves
            )
        ):
            leaves.append({
                **item,
                'type': 'Task',
                'id': f"{item['id']}-work",
                'title': item['title'],
                'parent_id': item['id']
            })

        if (
            item['type'] == 'Epic'
            and not any(
                x['parent_id'] == item['id']
                for x in staged
            )
        ):
            leaves.append({
                **item,
                'type': 'Task',
                'id': f"{item['id']}-work",
                'title': item['title'],
                'parent_id': item['id']
            })

    if not leaves:
        raise HTTPException(
            422,
            'SAD produced no actionable Story or Task'
        )

    # Catalog entries carry source-owned evidence,
    # not invented SAD references.
    epics, features, result = [], [], []

    def register(catalog, title, source):
        title = _clean_title(title)

        if not title:
            title = 'Architecture delivery'

        for existing in catalog:
            if _similarity(
                title,
                existing['title']
            ) >= .84:
                return existing

        entry = {
            'id': f'sad-{uuid.uuid4().hex[:12]}',
            'title': title,
            'source_section_id': source['source_section_id'],
            'source_excerpt': source['source_excerpt'],
            'description': '',
            'acceptance_criteria': [],
            'assumptions': [],
            'type': (
                'Epic'
                if catalog is epics
                else 'Feature'
            ),
            'parent_id': ''
        }

        catalog.append(entry)

        return entry

    for offset in range(0, len(leaves), 8):
        batch = leaves[
            offset:offset + 8
        ]

        inputs = []

        for leaf in batch:
            parent = by_id.get(
                leaf['parent_id'],
                {}
            )

            grandparent = by_id.get(
                parent.get('parent_id', ''),
                {}
            )

            inputs.append({
                'id': leaf['id'],
                'work': leaf['title'],
                'feature_hint': parent.get(
                    'title',
                    ''
                ),
                'epic_hint': (
                    grandparent.get('title', '')
                    or (
                        parent.get('title', '')
                        if parent.get('type') == 'Epic'
                        else ''
                    )
                ),
                'section': leaf[
                    'source_section_id'
                ]
            })

        catalog = [
            {
                'id': e['id'],
                'title': e['title']
            }
            for e in epics
        ]

        feature_catalog = [
            {
                'id': f['id'],
                'title': f['title'],
                'epic_id': f['parent_id']
            }
            for f in features
        ]

        assignments = {}

        try:
            answer = llm.generate_json(
                'Classify SAD work into a coherent DOCUMENT-WIDE backlog. '
                'Reuse existing epic and feature IDs where scope matches; '
                'do not create a new Epic/Feature per work item. '
                'An Epic is a broad product capability; '
                'a Feature groups related deliverables. '
                'Do not merge unrelated work. '
                'Return ONLY JSON '
                '{"assignments":[{"id":"input id","epic_id":"existing id or empty",'
                '"epic_title":"title if new","feature_id":"existing id or empty",'
                '"feature_title":"title if new"}]}. '
                'One assignment for EVERY input ID, no other fields. '
                'For a new feature under an existing epic, use epic_id. '
                'Keep names short.',
                json.dumps(
                    {
                        'work': inputs,
                        'epics': catalog,
                        'features': feature_catalog
                    },
                    ensure_ascii=False
                ),
                max_tokens=2400,
            )

            rows = answer.get(
                'assignments'
            )

            if (
                not isinstance(rows, list)
                or len(rows) != len(batch)
            ):
                raise ValueError(
                    'Classifier did not assign every work item'
                )

            assignments = {
                str(r.get('id')): r
                for r in rows
                if isinstance(r, dict)
            }

            if set(assignments) != {
                x['id']
                for x in batch
            }:
                raise ValueError(
                    'Classifier returned missing or duplicate IDs'
                )

        except (
            ValueError,
            TypeError,
            KeyError,
            RuntimeError
        ) as exc:

            logger.warning(
                'SAD hierarchy classification batch %s: %s; '
                'using conservative fallback',
                offset // 8 + 1,
                type(exc).__name__
            )

            assignments = {}

        for leaf, inp in zip(
            batch,
            inputs
        ):
            choice = assignments.get(
                leaf['id'],
                {}
            )

            epic = next(
                (
                    e
                    for e in epics
                    if e['id'] == choice.get('epic_id')
                ),
                None
            )

            if epic is None:
                proposed = _clean_title(
                    choice.get('epic_title')
                    or inp['epic_hint']
                    or 'Architecture delivery'
                )

                # Reuse closely matching broad capability titles,
                # never based on a shared generic word alone.
                epic = next(
                    (
                        e
                        for e in epics
                        if _similarity(
                            proposed,
                            e['title']
                        ) >= .65
                    ),
                    None
                )

                if epic is None:
                    epic = register(
                        epics,
                        proposed,
                        leaf
                    )

            feature = next(
                (
                    f
                    for f in features
                    if (
                        f['id']
                        == choice.get('feature_id')
                        and f['parent_id']
                        == epic['id']
                    )
                ),
                None
            )

            if feature is None:
                proposed = _clean_title(
                    choice.get('feature_title')
                    or inp['feature_hint']
                    or leaf['title']
                )

                feature = next(
                    (
                        f
                        for f in features
                        if (
                            f['parent_id']
                            == epic['id']
                            and _similarity(
                                proposed,
                                f['title']
                            ) >= .65
                        )
                    ),
                    None
                )

                if feature is None:
                    feature = register(
                        features,
                        proposed,
                        leaf
                    )

                    # register() may reuse a feature title from another epic;
                    # never cross-link it.
                    if (
                        feature.get('parent_id')
                        and feature['parent_id']
                        != epic['id']
                    ):
                        feature = {
                            **feature,
                            'id': f'sad-{uuid.uuid4().hex[:12]}'
                        }

                        features.append(feature)

                    feature['parent_id'] = epic['id']

            result.append({
                **leaf,
                'parent_id': feature['id']
            })

    output = epics + features + result

    ids = {
        item['id']
        for item in output
    }

    if len(ids) != len(output):
        raise HTTPException(
            502,
            'Hierarchy reconciliation produced duplicate IDs'
        )

    lookup = {
        item['id']: item
        for item in output
    }

    expected = {
        'Feature': 'Epic',
        'Story': 'Feature',
        'Task': 'Feature'
    }

    for item in output:
        if (
            item['type'] != 'Epic'
            and lookup.get(
                item['parent_id'],
                {}
            ).get('type')
            != expected[item['type']]
        ):
            raise HTTPException(
                502,
                'Hierarchy reconciliation produced an invalid parent link'
            )

    if len(result) != len(leaves):
        raise HTTPException(
            502,
            'Hierarchy reconciliation lost generated work'
        )

    return output


def _identify_clarifications(
    llm,
    text: str,
    document_title: str
):
    """Single pass over the whole SAD: surface genuine ambiguities before ticket generation.

    Always returns between 1 and 5 questions (enforced below), each with a suggested
    default_answer so the user can accept it as-is or edit it. Best-effort: if the LLM
    call fails outright, fall back to a single generic confirmation question rather than
    blocking proposal generation entirely.
    """

    fallback = [
        {
            'id': f'clarify-{uuid.uuid4().hex[:10]}',
            'question': (
                'Is this document ready to be broken down as written, '
                'or is anything missing?'
            ),
            'context': '',
            'default_answer': (
                'Proceed as written; no further clarification needed.'
            )
        }
    ]

    try:
        answer = llm.generate_json(
            'You are reviewing a Software Architecture Document before it '
            'is decomposed into a delivery backlog (Epic/Feature/Story/Task). '
            'Identify genuine ambiguities, missing information, undecided '
            'scope, or conflicting requirements that would materially change '
            'how the work should be broken down, sequenced, or estimated. '
            'Do not ask about writing style, formatting, or points already '
            'answered elsewhere in the document. If nothing is genuinely '
            'ambiguous, ask a single confirmation question about the most '
            'important assumption a delivery team would still be making. '
            'Return ONLY JSON {"questions": [{"question": "short specific '
            'question", "context": "short paraphrase of the relevant part '
            'of the document, not a verbatim quotation", "default_answer": '
            '"a reasonable default answer the user can accept as-is or edit"}]}. '
            'Return AT LEAST 1 and AT MOST 5 questions, highest-impact first.',
            json.dumps(
                {
                    'document_title': document_title,
                    'text': text
                },
                ensure_ascii=False
            ),
            max_tokens=1400,
        )

    except (
        ValueError,
        TypeError,
        KeyError,
        RuntimeError
    ) as exc:

        logger.warning(
            'SAD clarification pass failed: %s; '
            'using fallback confirmation question',
            type(exc).__name__
        )

        return fallback

    raw = (
        answer.get('questions')
        if isinstance(answer, dict)
        else None
    )

    if not isinstance(raw, list):
        return fallback

    questions = []

    for item in raw[:5]:
        if not isinstance(item, dict):
            continue

        question = sanitize_text(
            str(item.get('question', ''))
        ).text.strip()[:400]

        if not question:
            continue

        context = sanitize_text(
            str(item.get('context', ''))
        ).text.strip()[:300]

        default_answer = sanitize_text(
            str(item.get('default_answer', ''))
        ).text.strip()[:500]

        if not default_answer:
            default_answer = (
                'Proceed as written; no further clarification needed.'
            )

        questions.append({
            'id': f'clarify-{uuid.uuid4().hex[:10]}',
            'question': question,
            'context': context,
            'default_answer': default_answer
        })

    return (
        questions[:5]
        if questions
        else fallback
    )


def _apply_clarifications(
    text: str,
    clarifications
):
    """Fold user-answered clarification Q&A back into the document as a normal section.

    Reusing the existing section/chunk pipeline means answers flow into ticket
    generation the same way any other SAD content does, instead of needing a
    separate code path.
    """
    if not clarifications:
        return text

    lines = []

    for item in clarifications:
        if not isinstance(item, dict):
            continue

        question = sanitize_text(
            str(item.get('question', ''))
        ).text.strip()[:400]

        answer = sanitize_text(
            str(item.get('answer', ''))
        ).text.strip()[:1000]

        if question and answer:
            lines.append(
                f'Q: {question}\nA: {answer}'
            )

    if not lines:
        return text

    return (
        text.rstrip()
        + '\n\nClarifications\n'
        + '\n\n'.join(lines)
        + '\n'
    )


def _generate_chunks(llm, chunks):
    """Generate independent source chunks concurrently; return in document order.

    All-or-nothing: no proposal is returned if any chunk fails. A bounded worker
    count limits concurrent provider requests; set SAD_GENERATION_WORKERS=1 to
    restore the previous sequential behavior for rate-limited providers.
    """
    raw_workers = os.getenv(
        'SAD_GENERATION_WORKERS',
        '1'
    )

    try:
        workers = int(raw_workers)

    except ValueError as exc:
        raise HTTPException(
            500,
            'SAD_GENERATION_WORKERS must be an integer from 1 to 4'
        ) from exc

    if not 1 <= workers <= 4:
        raise HTTPException(
            500,
            'SAD_GENERATION_WORKERS must be an integer from 1 to 4'
        )

    started = time.monotonic()

    logger.info(
        'SAD generation started: chunks=%s workers=%s',
        len(chunks),
        min(workers, len(chunks))
    )

    if workers == 1 or len(chunks) == 1:
        results = [
            _generate_batch(
                llm,
                chunk,
                i,
                len(chunks)
            )
            for i, chunk in enumerate(
                chunks,
                start=1
            )
        ]

    else:
        results = [None] * len(chunks)

        with ThreadPoolExecutor(
            max_workers=min(
                workers,
                len(chunks)
            )
        ) as pool:

            futures = {
                pool.submit(
                    _generate_batch,
                    llm,
                    chunk,
                    i,
                    len(chunks)
                ): i - 1
                for i, chunk in enumerate(
                    chunks,
                    start=1
                )
            }

            for future in as_completed(futures):
                index = futures[future]

                try:
                    results[index] = future.result()

                except Exception:
                    for pending in futures:
                        pending.cancel()

                    raise

    logger.info(
        'SAD generation completed: chunks=%s elapsed_seconds=%.2f',
        len(chunks),
        time.monotonic() - started
    )

    return results


def generate_proposal(
    text: str,
    excel_path: Path,
    document_title: str,
    clarifications=None
):
    # Security boundary: sanitize SAD content before parsing/chunking
    # or any LLM call.
    sanitized_sad = sanitize_text(text)
    text = sanitized_sad.text.strip()

    # Prevent PII/secrets in the document title from reaching the response.
    sanitized_title = sanitize_text(document_title)
    document_title = sanitized_title.text.strip()

    if len(text) < 60:
        raise HTTPException(
            422,
            'Provide at least 60 characters of architecture content'
        )

    if len(text) > MAX_CHARS:
        raise HTTPException(
            413,
            f'SAD text must be {MAX_CHARS} characters or fewer'
        )

    llm = LLMClient()

    if not llm.enabled:
        raise HTTPException(
            503,
            'SAD generation requires LLM_PROVIDER=groq or bedrock. '
            'No tickets were generated.'
        )

    # First pass (no answers yet): ask the LLM whether it needs anything
    # clarified before committing to a ticket breakdown.
    # Skip this pass once answers come back.
    if not clarifications:
        questions = _identify_clarifications(
            llm,
            text,
            document_title
        )

        if questions:
            sections = sections_from_text(text)

            return {
                'status': 'needs_clarification',
                'document_title': document_title,
                'security': {
                    'input_sanitized': sanitized_sad.changed,
                    'categories_detected': sanitized_sad.categories,
                },
                'sections': [
                    {
                        'id': s['id'],
                        'title': s['title']
                    }
                    for s in sections
                ],
                'clarifications': questions,
                'tickets': [],
                'publish_status': 'awaiting_clarification',
            }

    else:
        text = _apply_clarifications(
            text,
            clarifications
        )

    sections = sections_from_text(text)
    chunks = _chunks(sections)

    if not chunks:
        raise HTTPException(
            422,
            'No readable SAD section content found'
        )

    # Stage each batch in memory.
    # On any error, return no proposal; never publish to Jira.
    staged = []

    for chunk, raw in zip(
        chunks,
        _generate_chunks(llm, chunks)
    ):

        if len(staged) + len(raw) > MAX_TICKETS:
            raise HTTPException(
                422,
                f'SAD would exceed {MAX_TICKETS} draft tickets; '
                'split document into smaller SADs. '
                'No partial proposal was returned.'
            )

        local_ids = {
            str(item['temp_id']).strip():
                f'sad-{uuid.uuid4().hex[:12]}'
            for item in raw
        }

        for item in raw:
            parent = str(
                item.get('parent_temp_id') or ''
            ).strip()

            criteria = item.get(
                'acceptance_criteria',
                []
            )

            dod = item.get(
                'definition_of_done',
                []
            )

            assumptions = item.get(
                'assumptions',
                []
            )

            staged.append({
                'id': local_ids[
                    str(item['temp_id']).strip()
                ],
                'parent_id': local_ids.get(
                    parent,
                    ''
                ),
                'type': item['type'],
                'title': str(
                    item['title']
                ).strip()[:255],
                'description': str(
                    item.get(
                        'description',
                        ''
                    )
                ).strip(),
                'acceptance_criteria': (
                    [
                        str(x)
                        for x in criteria
                    ]
                    if isinstance(
                        criteria,
                        list
                    )
                    else []
                ),
                'definition_of_done': (
                    [
                        str(x)
                        for x in dod
                        if str(x).strip()
                    ][:2]
                    if isinstance(
                        dod,
                        list
                    )
                    else []
                ),
                'source_section_id': chunk['id'],
                'source_excerpt': str(
                    item['source_excerpt']
                ).strip()[:500],
                'assumptions': (
                    [
                        str(x)
                        for x in assumptions
                    ]
                    if isinstance(
                        assumptions,
                        list
                    )
                    else []
                ),
            })

    if not staged:
        raise HTTPException(
            422,
            'No actionable requirements found in the SAD; '
            'no tickets were generated'
        )

    reconciliation_started = time.monotonic()

    staged = _reconcile_hierarchy(
        llm,
        staged
    )

    logger.info(
        'SAD hierarchy reconciliation completed: '
        'elapsed_seconds=%.2f',
        time.monotonic() - reconciliation_started
    )

    if len(staged) > MAX_TICKETS:
        raise HTTPException(
            422,
            f'Consolidated SAD exceeds {MAX_TICKETS} draft tickets; '
            'no partial proposal was returned'
        )

    try:
        backlog = pd.read_excel(
            excel_path,
            sheet_name='Backlog'
        ).fillna('')

        # Defense in depth:
        # ingest should already sanitize the workbook.
        backlog = sanitize_dataframe(
            backlog
        )

    except Exception as exc:
        logger.exception(
            'Unable to read SAD duplicate-check dataset'
        )

        raise HTTPException(
            422,
            'Unable to read the dataset Backlog sheet '
            'for duplicate screening'
        ) from exc

    for item in staged:
        item['duplicate_candidates'] = _candidate(
            item,
            backlog
        )

        item['review_status'] = (
            'potential_duplicate'
            if item['duplicate_candidates']
            else 'new_proposal'
        )

    return {
        'status': 'complete',
        'document_title': document_title,
        'security': {
            'input_sanitized': sanitized_sad.changed,
            'categories_detected': sanitized_sad.categories,
        },
        'sections': [
            {
                'id': s['id'],
                'title': s['title']
            }
            for s in sections
        ],
        'tickets': staged,
        'coverage': [
            {
                'section_id': s['id'],
                'title': s['title'],
                'ticket_count': sum(
                    t['source_section_id'] == s['id']
                    for t in staged
                )
            }
            for s in sections
        ],
        'duplicate_method': (
            'Conservative same-type title AND description screening; '
            'flagged matches require human review. '
            'Not semantic verification.'
        ),
        'clarifications_applied': bool(
            clarifications
        ),
        'publish_status': 'draft_only',
    }